# registry_service.py
"""Поиск, экспорт реестра и пакетная печать актов."""

import csv
import json
import logging
import re
import shutil
from pathlib import Path

from config import ACTS_DIR, ARCHIVE_DIR, resolve_path, to_stored_path
from db import get_acts_for_object, set_act_status, update_act
from storage import archive_act_to_archive
from utils import act_filename

logger = logging.getLogger(__name__)

EXPORT_HEADERS = [
    "№ акта",
    "Дата акта",
    "Работа",
    "Тип работы",
    "Наименование",
    "Позиция",
    "Раздел",
    "Шифр",
    "Листы",
    "Статус",
    "Материалы",
    "Файл",
    "Создан",
]

MONTHS = {
    "январ": 1,
    "феврал": 2,
    "март": 3,
    "апрел": 4,
    "май": 5,
    "мая": 5,
    "июн": 6,
    "июл": 7,
    "август": 8,
    "сентябр": 9,
    "октябр": 10,
    "ноябр": 11,
    "декабр": 12,
}

STOP_WORDS = {
    "за",
    "акт",
    "акта",
    "акте",
    "позиция",
    "позиции",
    "позицию",
    "поз",
    "поз.",
    "№",
    "на",
    "в",
    "от",
    "и",
    "по",
    "месяц",
}


def status_label(status: str | None) -> str:
    return {
        "draft": "Черновик",
        "ready": "Готовый",
        "printed": "Распечатан",
        "archived": "Архив",
    }.get(status or "ready", status or "—")


def _date_text(act: dict) -> str:
    return " ".join(str(act.get(field) or "") for field in ("act_date", "start_date", "end_date"))


def _material_text(act: dict) -> str:
    raw = act.get("materials_data") or ""
    try:
        items = json.loads(raw)
    except (TypeError, ValueError):
        return raw
    parts = []
    for item in items if isinstance(items, list) else []:
        if isinstance(item, dict):
            parts.append(str(item.get("name") or ""))
            if item.get("cert"):
                parts.append(str(item["cert"]))
    return " ".join(parts)


def _position_matches(position: str, query: str) -> bool:
    position = position.replace(" ", "")
    query = query.replace(" ", "")
    if position == query:
        return True
    match = re.fullmatch(r"(\d+)[-–](\d+)", position)
    if match and query.isdigit():
        return int(match.group(1)) <= int(query) <= int(match.group(2))
    return query in position


def search_acts(object_id: int, query: str, limit: int = 300):
    """Поиск по номеру, позиции, месяцу, дате, работе и материалам."""
    raw = (query or "").strip()
    if not raw:
        return []
    low = raw.lower().replace("ё", "е")

    number_match = re.search(r"(?:акт\w*|номер\w*|№)\s*№?\s*(\d{1,10})", low)
    number = number_match.group(1) if number_match else None
    position_match = re.search(r"(?:позици\w*|поз\.?)\s*№?\s*([0-9][0-9.,\-–]*)", low)
    position = position_match.group(1) if position_match else None
    exact_date_match = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", low)
    exact_date = None
    if exact_date_match:
        d, m, y = exact_date_match.groups()
        exact_date = f"{int(d):02d}.{int(m):02d}.{y}"
    month = None
    for stem, number_month in MONTHS.items():
        if re.search(rf"\b{stem}", low):
            month = number_month
            break
    year_match = re.search(r"\b(20\d{2})\b", low)
    year = year_match.group(1) if year_match else None
    if number is None and re.fullmatch(r"\d{1,10}", low):
        if len(low) == 4 and 2000 <= int(low) <= 2100:
            year = low
        else:
            number = low

    tokens = [
        token
        for token in re.findall(r"[0-9a-zа-я]+", low)
        if token
        and token not in STOP_WORDS
        and not token.isdigit()
        and not any(stem in token for stem in MONTHS)
    ]

    acts = get_acts_for_object(object_id, limit=100000)
    results = []
    for act in acts:
        if number is not None:
            act_number = str(act.get("act_number") or "").strip()
            if act_number.lstrip("0") != number.lstrip("0"):
                continue
        if position is not None and not _position_matches(
            str(act.get("position_number") or ""), position
        ):
            continue
        dates = _date_text(act)
        if exact_date and exact_date not in dates:
            continue
        if month is not None and f".{month:02d}." not in dates:
            continue
        if year and year not in dates:
            continue

        haystack = (
            " ".join(
                [
                    str(act.get("act_number") or ""),
                    str(act.get("work_full") or ""),
                    str(act.get("work_type_name") or ""),
                    str(act.get("work_name") or ""),
                    str(act.get("position_number") or ""),
                    str(act.get("section") or ""),
                    str(act.get("cipher") or ""),
                    dates,
                    _material_text(act),
                ]
            )
            .lower()
            .replace("ё", "е")
        )
        if tokens and not all(token in haystack for token in tokens):
            continue
        results.append(act)

    results.sort(
        key=lambda act: (
            int(act["act_number"]) if str(act.get("act_number") or "").isdigit() else 999999,
            act.get("created_at") or "",
        ),
        reverse=True,
    )
    return results[:limit]


def _export_row(act: dict) -> list:
    return [
        act.get("act_number") or "",
        act.get("act_date") or "",
        act.get("work_full") or "",
        act.get("work_type_name") or "",
        act.get("work_name") or "",
        act.get("position_number") or "",
        act.get("section") or "",
        act.get("cipher") or "",
        act.get("sheets") or "",
        status_label(act.get("status")),
        _material_text(act),
        act.get("file_path") or "",
        act.get("created_at") or "",
    ]


def acts_to_csv(acts, target: str | Path) -> Path:
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(EXPORT_HEADERS)
        for act in acts:
            writer.writerow(_export_row(act))
    return target


def acts_to_xlsx(acts, target: str | Path) -> Path:
    """Реестр актов в Excel с фильтром, закреплённой шапкой и шириной колонок."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реестр актов"
    sheet.append(EXPORT_HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for act in acts:
        sheet.append(_export_row(act))
    widths = [10, 16, 48, 20, 36, 12, 10, 16, 10, 14, 42, 40, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(target)
    return target


def merge_acts_to_docx(acts, target: str | Path) -> Path:
    """Склеивает готовые .docx в один файл с разрывом страницы между актами."""
    try:
        from docx import Document
        from docxcompose.composer import Composer
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError(
            "Для пакетной печати установите docxcompose: " "pip install docxcompose"
        ) from exc

    paths = []
    for act in acts:
        raw = act.get("file_path")
        if not raw:
            continue
        path = resolve_path(raw)
        if path.exists() and path.suffix.lower() == ".docx":
            paths.append(path)
    if not paths:
        raise RuntimeError("Нет доступных .docx для сборки пакета")

    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    master = Document(paths[0])
    composer = Composer(master)
    for path in paths[1:]:
        composer.doc.add_page_break()
        composer.append(Document(path))
    composer.save(target)
    return target


def rename_act_files(acts) -> tuple[int, int]:
    """Приводит имена существующих файлов к текущему формату."""
    renamed = skipped = 0
    for act in acts:
        raw = act.get("file_path")
        if not raw:
            skipped += 1
            continue
        source = resolve_path(raw)
        if not source.exists():
            skipped += 1
            continue
        target_dir = ACTS_DIR
        if (act.get("status") or "ready") == "archived":
            try:
                source.relative_to(ARCHIVE_DIR)
            except ValueError:
                skipped += 1
                continue
            target_dir = source.parent
        target = target_dir / act_filename(
            act.get("act_number") or "",
            act.get("work_type_name") or "",
            act.get("work_name") or "",
            act.get("position_number") or "",
        )
        if source.resolve() == target.resolve():
            skipped += 1
            continue
        counter = 1
        while target.exists():
            target = ACTS_DIR / (f"{target.stem}_{counter}{target.suffix}")
            counter += 1
        try:
            source.replace(target)
            update_act(act["id"], file_path=to_stored_path(target))
            renamed += 1
        except OSError:
            logger.exception("Не удалось переименовать акт %s", act.get("id"))
            skipped += 1
    return renamed, skipped


def archive_acts(acts) -> tuple[int, int]:
    """Переводит акты в архив; файлы переносит в acts/_archive/[год]/."""
    archived = failed = 0
    for act in acts:
        raw_date = str(act.get("act_date") or "")
        year_match = re.search(r"(20\d{2})", raw_date)
        year = year_match.group(1) if year_match else "unknown"
        source = resolve_path(act["file_path"]) if act.get("file_path") else None
        moved = None
        try:
            moved = archive_act_to_archive(act.get("file_path"), year)
            if source and source.exists() and not moved:
                raise OSError(f"не удалось перенести файл {source}")
            set_act_status(act["id"], "archived")
            archived += 1
        except Exception:
            if moved and moved.exists() and source and not source.exists():
                try:
                    source.parent.mkdir(parents=True, exist_ok=True)
                    shutil.move(str(moved), str(source))
                except OSError:
                    logger.exception("Не удалось вернуть файл акта %s", act.get("id"))
            failed += 1
            logger.exception("Не удалось архивировать акт %s", act.get("id"))
    return archived, failed
